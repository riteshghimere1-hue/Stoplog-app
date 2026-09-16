from flask import Flask, request, jsonify, send_file, send_from_directory
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
import sqlite3
import os

app = Flask(__name__, static_folder='public', static_url_path='')

# --- DATABASE SETUP ---
# On your PC, this saves next to app.py. 
# When you deploy, change this path to a persistent folder.
DB_PATH = os.environ.get('DATABASE_PATH', 'database.sqlite')


def get_db():
    """Open a connection to the SQLite database."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row  # Lets us access columns by name
    return conn


def init_db():
    """Create tables if they don't exist yet."""
    conn = get_db()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS canals (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS stoplogs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        canal_id INTEGER,
        location TEXT,
        L REAL,
        H REAL,
        T REAL,
        No INTEGER,
        total_height REAL,
        qty_sqm REAL,
        qty_cum REAL,
        FOREIGN KEY(canal_id) REFERENCES canals(id)
    )''')
    # Add a default canal if the table is empty
    c.execute('SELECT COUNT(*) FROM canals')
    if c.fetchone()[0] == 0:
        c.execute("INSERT INTO canals (name) VALUES ('Beluwa SSC')")
    conn.commit()
    conn.close()


# --- STATIC FILE SERVING ---
@app.route('/')
def index():
    return send_from_directory('public', 'index.html')


# --- API: CANALS ---
@app.route('/api/canals', methods=['GET'])
def get_canals():
    conn = get_db()
    rows = conn.execute('SELECT * FROM canals ORDER BY name ASC').fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route('/api/canals', methods=['POST'])
def add_canal():
    name = request.json.get('name')
    try:
        conn = get_db()
        c = conn.cursor()
        c.execute('INSERT INTO canals (name) VALUES (?)', (name,))
        conn.commit()
        new_id = c.lastrowid
        conn.close()
        return jsonify({'id': new_id, 'name': name})
    except sqlite3.IntegrityError:
        return jsonify({'error': 'Canal already exists'}), 400


@app.route('/api/canals/<int:canal_id>', methods=['PUT'])
def update_canal(canal_id):
    name = request.json.get('name')
    try:
        conn = get_db()
        conn.execute('UPDATE canals SET name = ? WHERE id = ?', (name, canal_id))
        conn.commit()
        conn.close()
        return jsonify({'updated': 1})
    except sqlite3.IntegrityError:
        return jsonify({'error': 'Canal name already exists'}), 400


# --- API: STOPLOGS ---
@app.route('/api/stoplogs', methods=['GET'])
def get_stoplogs():
    conn = get_db()
    query = '''
        SELECT s.*, c.name as canal_name 
        FROM stoplogs s 
        JOIN canals c ON s.canal_id = c.id 
        ORDER BY c.name ASC, s.id ASC
    '''
    rows = conn.execute(query).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route('/api/stoplogs', methods=['POST'])
def add_stoplog():
    data = request.json
    L, H, T, No = data['L'], data['H'], data['T'], data['No']
    total_height = H * No
    qty_sqm = (L * total_height) / 1_000_000
    qty_cum = (L * T * total_height) / 1_000_000_000

    conn = get_db()
    c = conn.cursor()
    c.execute('''INSERT INTO stoplogs 
        (canal_id, location, L, H, T, No, total_height, qty_sqm, qty_cum) 
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
        (data['canal_id'], data['location'], L, H, T, No, total_height, qty_sqm, qty_cum))
    conn.commit()
    new_id = c.lastrowid
    conn.close()
    return jsonify({'id': new_id})


@app.route('/api/stoplogs/<int:entry_id>', methods=['PUT'])
def update_stoplog(entry_id):
    data = request.json
    L, H, T, No = data['L'], data['H'], data['T'], data['No']
    total_height = H * No
    qty_sqm = (L * total_height) / 1_000_000
    qty_cum = (L * T * total_height) / 1_000_000_000

    conn = get_db()
    conn.execute('''UPDATE stoplogs 
        SET canal_id = ?, location = ?, L = ?, H = ?, T = ?, No = ?, 
            total_height = ?, qty_sqm = ?, qty_cum = ? 
        WHERE id = ?''',
        (data['canal_id'], data['location'], L, H, T, No,
         total_height, qty_sqm, qty_cum, entry_id))
    conn.commit()
    conn.close()
    return jsonify({'updated': 1})


@app.route('/api/stoplogs/<int:entry_id>', methods=['DELETE'])
def delete_stoplog(entry_id):
    conn = get_db()
    conn.execute('DELETE FROM stoplogs WHERE id = ?', (entry_id,))
    conn.commit()
    conn.close()
    return jsonify({'deleted': 1})


# --- EXCEL EXPORT ---
@app.route('/api/export')
def export_excel():
    conn = get_db()
    rows = conn.execute('''
        SELECT s.*, c.name as canal_name 
        FROM stoplogs s 
        JOIN canals c ON s.canal_id = c.id 
        ORDER BY c.name ASC, s.id ASC
    ''').fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Stoplog Quantity"

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Row 1: Title
    ws.merge_cells('A1:H1')
    title = ws['A1']
    title.value = 'Stoplog Quantity'
    title.font = Font(bold=True, size=16)
    title.alignment = Alignment(horizontal='center')

    # Row 2: Headers
    headers = ['Location', 'L (mm)', 'H (mm)', 'T (mm)', 'No',
               'Total Height (mm)', 'Quantity (m²)', 'Quantity (m³)']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    current_row = 3
    current_canal = ''
    total_sqm = 0.0
    total_cum = 0.0

    for row in rows:
        # Canal header row when the canal name changes
        if row['canal_name'] != current_canal:
            current_canal = row['canal_name']
            ws.merge_cells(start_row=current_row, start_column=1,
                           end_row=current_row, end_column=8)
            cell = ws.cell(row=current_row, column=1, value=current_canal)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='FFE0E0E0',
                                    end_color='FFE0E0E0', fill_type='solid')
            current_row += 1

        # Data row
        ws.cell(row=current_row, column=1, value=row['location'])
        ws.cell(row=current_row, column=2, value=row['L'])
        ws.cell(row=current_row, column=3, value=row['H'])
        ws.cell(row=current_row, column=4, value=row['T'])
        ws.cell(row=current_row, column=5, value=row['No'])
        ws.cell(row=current_row, column=6, value=row['total_height'])
        ws.cell(row=current_row, column=7, value=row['qty_sqm']).number_format = '0.00'
        ws.cell(row=current_row, column=8, value=row['qty_cum']).number_format = '0.0000'

        total_sqm += row['qty_sqm']
        total_cum += row['qty_cum']
        current_row += 1

    # Total row
    ws.cell(row=current_row, column=1, value='Total Quantity').font = Font(bold=True)
    c7 = ws.cell(row=current_row, column=7, value=total_sqm)
    c8 = ws.cell(row=current_row, column=8, value=total_cum)
    c7.font = Font(bold=True)
    c8.font = Font(bold=True)
    c7.number_format = '0.00'
    c8.number_format = '0.0000'

    # Borders everywhere
    for r in range(1, current_row + 1):
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = border

    # Column widths
    widths = [25, 12, 12, 12, 8, 18, 15, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Stoplog_Quantity.xlsx'
    )

init_db()

# --- START THE APP ---
if __name__ == '__main__':
    
    port = int(os.environ.get('PORT', 5000))
    print(f"\n🚀 App running! Open http://localhost:{port} in your browser.\n")
    app.run(host='0.0.0.0', port=port, debug=True)